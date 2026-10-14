# -*- coding: utf-8 -*-
"""
开题答辩 PPT v2 - 精致版
设计原则：大图主导、数据突出、卡片式排版、文字精简
"""
import os
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn

ROOT = r"D:\programme file\ResearchArticle"
RESULT_DIR = os.path.join(ROOT, "results")
OUT_PATH = os.path.join(ROOT, "开题答辩_双链路智能质检辅助系统.pptx")

# ===== 配色方案 =====
C_NAVY    = RGBColor(0x0A, 0x1F, 0x44)   # 深海军蓝（封面/深色页）
C_BLUE    = RGBColor(0x15, 0x65, 0xC0)   # 主蓝
C_BLUE_L  = RGBColor(0x42, 0xA5, 0xF5)   # 亮蓝
C_BLUE_BG = RGBColor(0xE3, 0xF2, 0xFD)   # 浅蓝背景
C_ORANGE  = RGBColor(0xE6, 0x51, 0x00)   # 强调橙
C_ORANGE_BG = RGBColor(0xFF, 0xF3, 0xE0) # 浅橙背景
C_GREEN   = RGBColor(0x2E, 0x7D, 0x32)   # 绿
C_GREEN_BG= RGBColor(0xE8, 0xF5, 0xE9)   # 浅绿背景
C_RED     = RGBColor(0xC6, 0x28, 0x28)   # 红
C_RED_BG  = RGBColor(0xFF, 0xEB, 0xEE)   # 浅红背景
C_GRAY    = RGBColor(0x55, 0x55, 0x55)   # 正文灰
C_GRAY_L  = RGBColor(0x9E, 0x9E, 0x9E)   # 浅灰
C_BG      = RGBColor(0xFA, 0xFA, 0xFA)   # 页面背景
C_WHITE   = RGBColor(0xFF, 0xFF, 0xFF)
C_DARK    = RGBColor(0x21, 0x21, 0x21)
C_LINE    = RGBColor(0xE0, 0xE0, 0xE0)   # 分割线

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
SW = prs.slide_width
SH = prs.slide_height

def blank():
    return prs.slides.add_slide(prs.slide_layouts[6])

def set_bg(slide, color):
    bg = slide.background
    fill = bg.fill
    fill.solid()
    fill.fore_color.rgb = color

def rect(slide, l, t, w, h, fill, line=None, line_w=1):
    s = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, l, t, w, h)
    s.fill.solid(); s.fill.fore_color.rgb = fill
    if line:
        s.line.color.rgb = line; s.line.width = Pt(line_w)
    else:
        s.line.fill.background()
    return s

def rrect(slide, l, t, w, h, fill, line=None, line_w=1.5):
    s = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, l, t, w, h)
    s.fill.solid(); s.fill.fore_color.rgb = fill
    if line:
        s.line.color.rgb = line; s.line.width = Pt(line_w)
    else:
        s.line.fill.background()
    return s

def txt(slide, l, t, w, h, text, size=18, bold=False, color=C_DARK,
        align=PP_ALIGN.LEFT, font="微软雅黑", anchor=MSO_ANCHOR.TOP):
    tb = slide.shapes.add_textbox(l, t, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = Pt(2); tf.margin_right = Pt(2)
    tf.margin_top = Pt(1); tf.margin_bottom = Pt(1)
    p = tf.paragraphs[0]
    p.text = text
    p.font.size = Pt(size); p.font.bold = bold
    p.font.color.rgb = color; p.font.name = font
    p.alignment = align
    return tb

def multiline(slide, l, t, w, h, lines, size=14, color=C_GRAY, spacing=1.4, font="微软雅黑"):
    tb = slide.shapes.add_textbox(l, t, w, h)
    tf = tb.text_frame; tf.word_wrap = True
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        if isinstance(line, tuple):
            text, b, c = line[0], line[1] if len(line)>1 else False, line[2] if len(line)>2 else color
        else:
            text, b, c = line, False, color
        p.text = text
        p.font.size = Pt(size); p.font.bold = b
        p.font.color.rgb = c; p.font.name = font
        p.space_after = Pt(size * 0.25)
        p.line_spacing = spacing
    return tb

def img(slide, path, l, t, width=None, height=None):
    if os.path.exists(path):
        if width and height: slide.shapes.add_picture(path, l, t, width=width, height=height)
        elif width: slide.shapes.add_picture(path, l, t, width=width)
        elif height: slide.shapes.add_picture(path, l, t, height=height)
        else: slide.shapes.add_picture(path, l, t)
        return True
    return False

def page_header(slide, title, subtitle=None, num=None):
    """统一页面顶部：左侧标题 + 底部细线"""
    rect(slide, 0, 0, SW, Inches(0.08), C_BLUE)
    txt(slide, Inches(0.6), Inches(0.25), Inches(10), Inches(0.55),
        title, size=26, bold=True, color=C_NAVY)
    if subtitle:
        txt(slide, Inches(0.6), Inches(0.78), Inches(10), Inches(0.35),
            subtitle, size=13, color=C_GRAY_L)
    if num:
        txt(slide, Inches(12.3), Inches(0.3), Inches(0.7), Inches(0.4),
            f"{num:02d}", size=16, bold=True, color=C_BLUE_L, align=PP_ALIGN.RIGHT)
    rect(slide, Inches(0.6), Inches(1.15), Inches(12.1), Pt(1.5), C_LINE)

def big_number(slide, l, t, number, label, color=C_BLUE, num_size=60):
    """超大数字 + 标签"""
    txt(slide, l, t, Inches(3), Inches(0.9), number, size=num_size, bold=True,
        color=color, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    txt(slide, l, t + Inches(0.85), Inches(3), Inches(0.35), label, size=12,
        color=C_GRAY, align=PP_ALIGN.CENTER)

# ================================================================
# 第1页：封面（深色大气）
# ================================================================
s = blank()
set_bg(s, C_NAVY)
# 装饰：右侧渐变条
rect(s, Inches(11.5), 0, Inches(1.833), SH, RGBColor(0x0D, 0x47, 0xA1))
rect(s, Inches(11.3), 0, Inches(0.15), SH, C_ORANGE)
# 顶部标签
txt(s, Inches(1.0), Inches(1.2), Inches(8), Inches(0.4),
    "《人工智能应用实践》课程开题报告", size=16, color=C_BLUE_L)
# 主标题
txt(s, Inches(1.0), Inches(2.2), Inches(10), Inches(1.0),
    "面向金属表面缺陷的", size=34, bold=True, color=C_WHITE)
txt(s, Inches(1.0), Inches(3.1), Inches(10), Inches(1.2),
    "双链路门控智能质检辅助系统", size=44, bold=True, color=C_WHITE)
# 橙色分隔线
rect(s, Inches(1.0), Inches(4.5), Inches(2.5), Pt(3), C_ORANGE)
# 一句话定位
txt(s, Inches(1.0), Inches(4.8), Inches(9), Inches(0.5),
    "置信度门控 · 少样本换产 · 引用可溯源", size=20, color=C_BLUE_L)
# 底部信息
txt(s, Inches(1.0), Inches(6.2), Inches(6), Inches(0.35),
    "山东大学人工智能学院", size=14, color=C_GRAY_L)
txt(s, Inches(1.0), Inches(6.55), Inches(8), Inches(0.35),
    "团队：XXX / XXX / XXX    指导教师：赵玉琨 副教授    2026.09",
    size=12, color=C_GRAY_L)

# ================================================================
# 第2页：目录
# ================================================================
s = blank()
set_bg(s, C_BG)
page_header(s, "目录", "CONTENTS", 2)

items = [
    ("01", "问题与需求", "行业痛点 · 系统定位"),
    ("02", "研究现状", "技术路线 · 现有不足"),
    ("03", "系统架构", "双链路门控 · 可溯源校验"),
    ("04", "创新点", "架构 · 工作流 · 工程实现"),
    ("05", "可行性与先导实验", "技术验证 · 实测数据"),
    ("06", "评估与计划", "指标体系 · 进度分工"),
]
for i, (num, title, desc) in enumerate(items):
    row = i // 2; col = i % 2
    l = Inches(0.8 + col * 6.1)
    t = Inches(1.6 + row * 1.7)
    rrect(s, l, t, Inches(5.7), Inches(1.4), C_WHITE, C_LINE, 1)
    # 序号块
    rect(s, l, t, Inches(1.1), Inches(1.4), C_BLUE)
    txt(s, l, t + Inches(0.25), Inches(1.1), Inches(0.9), num, size=36, bold=True,
        color=C_WHITE, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    txt(s, l + Inches(1.3), t + Inches(0.25), Inches(4.2), Inches(0.5),
        title, size=20, bold=True, color=C_NAVY)
    txt(s, l + Inches(1.3), t + Inches(0.8), Inches(4.2), Inches(0.4),
        desc, size=13, color=C_GRAY)

# ================================================================
# 第3页：行业痛点（3大卡片）
# ================================================================
s = blank()
set_bg(s, C_BG)
page_header(s, "一、待解决的问题", "中小机械加工厂金属零件多品种小批量离线抽检", 3)

pains = [
    ("换产成本高", "新零件需标注缺陷、训练模型", "周期长 · 成本高 · 小批量难承担", C_RED, C_RED_BG),
    ("分析效率低", "翻标准、写报告、分析根因", "依赖人工经验 · 单份报告30分钟+", C_ORANGE, C_ORANGE_BG),
    ("合规性弱", "人工判定标准不统一", "无溯源记录 · 不符合质检可追溯要求", C_BLUE, C_BLUE_BG),
]
for i, (title, sub, desc, c, bg) in enumerate(pains):
    l = Inches(0.6 + i * 4.15)
    t = Inches(1.6)
    rrect(s, l, t, Inches(3.85), Inches(4.2), C_WHITE, C_LINE, 1)
    # 顶部色条
    rect(s, l, t, Inches(3.85), Inches(0.12), c)
    # 大序号
    txt(s, l + Inches(0.3), t + Inches(0.4), Inches(1), Inches(0.8),
        f"0{i+1}", size=48, bold=True, color=bg)
    # 标题
    txt(s, l + Inches(0.3), t + Inches(1.5), Inches(3.3), Inches(0.5),
        title, size=24, bold=True, color=c)
    # 分隔
    rect(s, l + Inches(0.3), t + Inches(2.15), Inches(1.2), Pt(2), c)
    # 描述
    txt(s, l + Inches(0.3), t + Inches(2.4), Inches(3.3), Inches(0.5),
        sub, size=14, color=C_DARK)
    txt(s, l + Inches(0.3), t + Inches(3.0), Inches(3.3), Inches(0.8),
        desc, size=13, color=C_GRAY)

# 底部结论
rrect(s, Inches(0.6), Inches(6.2), Inches(12.1), Inches(0.8), C_NAVY)
txt(s, Inches(0.6), Inches(6.3), Inches(12.1), Inches(0.6),
    "现有方案共性不足：纯检测无分析 · 纯大模型无定位 · 定制系统成本高 — 缺少兼顾效率、精度、合规的完整方案",
    size=15, bold=True, color=C_WHITE, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)

# ================================================================
# 第4页：系统定位
# ================================================================
s = blank()
set_bg(s, C_BG)
page_header(s, "一、系统定位与目标", "人在回路的离线质检辅助工具", 4)

# 定位大卡片
rrect(s, Inches(0.6), Inches(1.5), Inches(12.1), Inches(1.6), C_NAVY)
txt(s, Inches(1.0), Inches(1.7), Inches(2), Inches(0.4),
    "系统定位", size=16, bold=True, color=C_ORANGE)
txt(s, Inches(1.0), Inches(2.15), Inches(11), Inches(0.8),
    "AI负责定位、初筛、整理依据，人工负责最终确认 — 不替代人工，做质检的「副驾驶」",
    size=22, bold=True, color=C_WHITE)

# 三个目标
goals = [
    ("少样本检测", "仅需10-20张正常样本建库\n换产无需模型训练", C_BLUE, "01"),
    ("自动化分析", "自动匹配标准、生成报告\n降低质检人工工作量", C_GREEN, "02"),
    ("可溯源合规", "所有结论绑定标准原文页码\n错误可识别、可校验", C_ORANGE, "03"),
]
for i, (title, desc, c, num) in enumerate(goals):
    l = Inches(0.6 + i * 4.15)
    t = Inches(3.5)
    rrect(s, l, t, Inches(3.85), Inches(2.8), C_WHITE, c, 2)
    rect(s, l, t, Inches(3.85), Inches(0.7), c)
    txt(s, l + Inches(0.25), t + Inches(0.12), Inches(1), Inches(0.45),
        num, size=22, bold=True, color=C_WHITE)
    txt(s, l + Inches(1.2), t + Inches(0.15), Inches(2.5), Inches(0.45),
        title, size=20, bold=True, color=C_WHITE)
    txt(s, l + Inches(0.3), t + Inches(1.1), Inches(3.3), Inches(1.5),
        desc, size=15, color=C_DARK)

# 应用场景
txt(s, Inches(0.6), Inches(6.6), Inches(12.1), Inches(0.4),
    "应用场景：金属零件表面缺陷离线抽检  ·  新产线快速质检搭建  ·  质检新人培训辅助",
    size=14, color=C_GRAY, align=PP_ALIGN.CENTER)

# ================================================================
# 第5页：研究现状
# ================================================================
s = blank()
set_bg(s, C_BG)
page_header(s, "二、研究现状", "工业缺陷检测两条技术路线 + 大模型RAG应用", 5)

# 左：两条技术路线
rrect(s, Inches(0.6), Inches(1.5), Inches(5.9), Inches(2.3), C_WHITE, C_RED, 1.5)
rect(s, Inches(0.6), Inches(1.5), Inches(5.9), Inches(0.55), C_RED)
txt(s, Inches(0.8), Inches(1.58), Inches(5.5), Inches(0.4),
    "监督学习路线（YOLO / CNN）", size=16, bold=True, color=C_WHITE)
multiline(s, Inches(0.85), Inches(2.25), Inches(5.5), Inches(1.4),
    ["✓ 检测精度高", "✗ 需大量标注数据，训练周期长", "✗ 换产重新标注成本高，少样本适配差"],
    size=13, color=C_DARK)

rrect(s, Inches(0.6), Inches(4.0), Inches(5.9), Inches(2.3), C_WHITE, C_GREEN, 1.5)
rect(s, Inches(0.6), Inches(4.0), Inches(5.9), Inches(0.55), C_GREEN)
txt(s, Inches(0.8), Inches(4.08), Inches(5.5), Inches(0.4),
    "零样本异常检测（WinCLIP等）", size=16, bold=True, color=C_WHITE)
multiline(s, Inches(0.85), Inches(4.75), Inches(5.5), Inches(1.4),
    ["✓ 仅需正常样本建库，适配小批量换产", "✓ 无需标注缺陷数据", "✗ 只有检测定位，缺少深度分析能力"],
    size=13, color=C_DARK)

# 右：大模型RAG + 不足
rrect(s, Inches(6.8), Inches(1.5), Inches(5.9), Inches(2.3), C_WHITE, RGBColor(0x6A,0x1B,0x9A), 1.5)
rect(s, Inches(6.8), Inches(1.5), Inches(5.9), Inches(0.55), RGBColor(0x6A,0x1B,0x9A))
txt(s, Inches(7.0), Inches(1.58), Inches(5.5), Inches(0.4),
    "生成式VLM + RAG知识库", size=16, bold=True, color=C_WHITE)
multiline(s, Inches(7.05), Inches(2.25), Inches(5.5), Inches(1.4),
    ["✓ 能理解图像、生成描述、结合知识库问答", "✗ 多为纯文本方案，缺少与视觉检测融合", "✗ 幻觉不可控，不符合工业合规要求"],
    size=13, color=C_DARK)

rrect(s, Inches(6.8), Inches(4.0), Inches(5.9), Inches(2.3), C_NAVY)
txt(s, Inches(7.0), Inches(4.2), Inches(5.5), Inches(0.4),
    "现有方案共性不足", size=16, bold=True, color=C_ORANGE)
multiline(s, Inches(7.05), Inches(4.75), Inches(5.5), Inches(1.4),
    ["要么只有检测没有分析", "要么只有分析没有定位", "缺少兼顾效率、精度、合规性的完整工作流"],
    size=14, color=C_WHITE, spacing=1.6)

# ================================================================
# 第6页：双链路门控机制（大图主导）
# ================================================================
s = blank()
set_bg(s, C_BG)
page_header(s, "三、核心机制 · 双链路门控", "高置信快检 · 低置信精析 · 置信度智能路由", 6)

# 左侧：机制要点（窄）
points = [
    ("感知层统一计算", "WinCLIP像素级异常分数 → 全局置信度", C_BLUE),
    ("高置信 → 快链路", "分数≥阈值，直接输出位置/等级/热力图", C_GREEN),
    ("低置信 → 慢链路", "分数<阈值，裁剪→VLM描述→RAG匹配→报告", C_ORANGE),
    ("三级反馈反哺", "案例库→阈值辅助调整→标准库人工审核", RGBColor(0x6A,0x1B,0x9A)),
]
for i, (title, desc, c) in enumerate(points):
    t = Inches(1.4 + i * 1.35)
    rrect(s, Inches(0.6), t, Inches(4.0), Inches(1.15), C_WHITE, c, 1.5)
    rect(s, Inches(0.6), t, Inches(0.1), Inches(1.15), c)
    txt(s, Inches(0.85), t + Inches(0.1), Inches(3.6), Inches(0.4),
        title, size=15, bold=True, color=c)
    txt(s, Inches(0.85), t + Inches(0.55), Inches(3.6), Inches(0.55),
        desc, size=11, color=C_GRAY)

# 阈值标注
rrect(s, Inches(0.6), Inches(6.85), Inches(4.0), Inches(0.45), C_ORANGE_BG)
txt(s, Inches(0.7), Inches(6.88), Inches(3.8), Inches(0.4),
    "阈值：中位数+1.5×MAD = 0.2243", size=12, bold=True, color=C_ORANGE,
    anchor=MSO_ANCHOR.MIDDLE)

# 右侧：双链路演示大图
img(s, os.path.join(RESULT_DIR, "gating", "dual_path_demo.png"),
    Inches(4.85), Inches(1.35), width=Inches(8.2))

# ================================================================
# 第7页：四层一底座架构
# ================================================================
s = blank()
set_bg(s, C_BG)
page_header(s, "三、系统架构 · 四层一底座", "数据输入 → 感知检测 → 门控路由 → 认知分析 → 应用交互", 7)

layers = [
    ("应用交互层", "单图检测 | 批量处理 | 报告导出 | 知识库管理 | 历史记录", RGBColor(0x6A,0x1B,0x9A)),
    ("认知分析层", "VLM缺陷描述 | RAG标准检索 | 根因推理 | 结构化报告", C_GREEN),
    ("门控路由层", "置信度计算 | 阈值判断 | 链路调度 | 冲突处理", C_ORANGE),
    ("感知检测层", "WinCLIP特征库 | 多尺度滑窗 | 异常热力图 | 缺陷裁剪", C_BLUE),
    ("数据输入层", "正常样本图 | 待检工件图 | 质检标准文档 | 历史案例", C_GRAY),
]
for i, (name, content, c) in enumerate(layers):
    t = Inches(1.4 + i * 0.95)
    rrect(s, Inches(1.5), t, Inches(10.3), Inches(0.8), C_WHITE, c, 1.5)
    rect(s, Inches(1.5), t, Inches(2.2), Inches(0.8), c)
    txt(s, Inches(1.5), t + Inches(0.18), Inches(2.2), Inches(0.45),
        name, size=17, bold=True, color=C_WHITE, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    txt(s, Inches(3.9), t + Inches(0.2), Inches(7.7), Inches(0.45),
        content, size=14, color=C_DARK, anchor=MSO_ANCHOR.MIDDLE)

# 底座
rrect(s, Inches(1.5), Inches(6.3), Inches(10.3), Inches(0.65), C_NAVY)
txt(s, Inches(1.5), Inches(6.38), Inches(10.3), Inches(0.5),
    "知识库底座：正常特征库 | 质检标准库 | 历史案例库 | 报告模板库",
    size=15, bold=True, color=C_WHITE, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)

# ================================================================
# 第8页：可溯源校验（对比大图）
# ================================================================
s = blank()
set_bg(s, C_BG)
page_header(s, "三、核心机制 · 页码级可溯源校验", "不追求零幻觉，追求所有引用可校验、错误可识别", 8)

# 左侧：5步机制
steps = [
    "片段ID化：每段标准绑定原文+页码+标准号",
    "RAG检索：返回片段+引用ID+页码",
    "生成约束：Prompt强制VLM输出携带引用ID",
    "后处理校验：自动检查ID是否在检索结果中",
    "降级模式：校验失败→展示原始片段供人工选",
]
for i, step in enumerate(steps):
    t = Inches(1.4 + i * 1.05)
    rrect(s, Inches(0.6), t, Inches(3.8), Inches(0.9), C_WHITE, C_GREEN, 1)
    # 圆形序号
    circ = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(0.75), t + Inches(0.2), Inches(0.5), Inches(0.5))
    circ.fill.solid(); circ.fill.fore_color.rgb = C_GREEN
    circ.line.fill.background()
    txt(s, Inches(0.75), t + Inches(0.22), Inches(0.5), Inches(0.45),
        str(i+1), size=16, bold=True, color=C_WHITE, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    txt(s, Inches(1.4), t + Inches(0.15), Inches(2.9), Inches(0.65),
        step, size=11, color=C_DARK, anchor=MSO_ANCHOR.MIDDLE)

# 右侧：通过 vs 失败 对比大图
txt(s, Inches(4.7), Inches(1.3), Inches(4), Inches(0.35),
    "✓ 校验通过", size=16, bold=True, color=C_GREEN)
img(s, os.path.join(RESULT_DIR, "cognition", "ref_validation_pass.png"),
    Inches(4.7), Inches(1.7), width=Inches(8.2))

txt(s, Inches(4.7), Inches(4.35), Inches(4), Inches(0.35),
    "✗ 编造引用 → 自动标红降级", size=16, bold=True, color=C_RED)
img(s, os.path.join(RESULT_DIR, "cognition", "ref_validation_fail.png"),
    Inches(4.7), Inches(4.75), width=Inches(8.2))

# ================================================================
# 第9页：创新点
# ================================================================
s = blank()
set_bg(s, C_BG)
page_header(s, "四、项目创新点", "底层模型用开源，创新集中在应用架构与工程落地", 9)

innovations = [
    ("应用架构创新", "置信度门控的双链路质检架构",
     "高置信快检、低置信精析，兼顾效率与深度；避免全量图走大模型的算力浪费",
     C_BLUE, C_BLUE_BG),
    ("场景工作流创新", "少样本全流程质检闭环",
     "建库→检测→定位→分析→报告→迭代；仅需10-20张正常样本，换产0.5-1天",
     C_GREEN, C_GREEN_BG),
    ("工程实现创新", "引用ID绑定的可溯源校验机制",
     "所有判定绑定标准原文页码，输出自动校验；失败自动降级，符合工业合规",
     C_ORANGE, C_ORANGE_BG),
]
for i, (cat, title, desc, c, bg) in enumerate(innovations):
    l = Inches(0.6 + i * 4.15)
    t = Inches(1.5)
    rrect(s, l, t, Inches(3.85), Inches(4.5), C_WHITE, c, 2)
    # 顶部色块
    rect(s, l, t, Inches(3.85), Inches(1.0), c)
    txt(s, l + Inches(0.25), t + Inches(0.15), Inches(3.4), Inches(0.35),
        cat, size=14, bold=True, color=bg)
    txt(s, l + Inches(0.25), t + Inches(0.5), Inches(3.4), Inches(0.45),
        title, size=18, bold=True, color=C_WHITE)
    # 大序号水印
    txt(s, l + Inches(0.2), t + Inches(1.2), Inches(2), Inches(1.0),
        f"0{i+1}", size=64, bold=True, color=bg)
    # 描述
    txt(s, l + Inches(0.3), t + Inches(2.5), Inches(3.3), Inches(1.8),
        desc, size=14, color=C_DARK)

# 底部技术边界
rrect(s, Inches(0.6), Inches(6.3), Inches(12.1), Inches(0.7), C_NAVY)
txt(s, Inches(0.6), Inches(6.38), Inches(12.1), Inches(0.55),
    "技术边界：底层视觉模型、语言模型、RAG框架均采用开源成熟组件，不做模型层原创",
    size=14, bold=True, color=C_ORANGE, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)

# ================================================================
# 第10页：技术可行性（超大数字 + 分数分布图）
# ================================================================
s = blank()
set_bg(s, C_BG)
page_header(s, "五、可行性分析 · 技术与数据", "MVTec-AD metal_nut 零样本无训练实测", 10)

# 左侧：4个超大数字指标
metrics = [
    ("0.9868", "检测 AUROC", C_BLUE),
    ("94.78%", "准确率", C_GREEN),
    ("96.77%", "精确率 / 召回率", C_ORANGE),
    ("0.2243", "门控阈值", RGBColor(0x6A,0x1B,0x9A)),
]
for i, (num, label, c) in enumerate(metrics):
    row = i // 2; col = i % 2
    l = Inches(0.6 + col * 2.9)
    t = Inches(1.5 + row * 2.3)
    rrect(s, l, t, Inches(2.7), Inches(2.0), C_WHITE, c, 2)
    txt(s, l, t + Inches(0.3), Inches(2.7), Inches(1.0), num, size=40, bold=True,
        color=c, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    txt(s, l, t + Inches(1.35), Inches(2.7), Inches(0.4), label, size=13,
        color=C_GRAY, align=PP_ALIGN.CENTER)

# 右侧：分数分布图 + 说明
rrect(s, Inches(6.6), Inches(1.5), Inches(6.2), Inches(4.6), C_WHITE, C_LINE, 1)
txt(s, Inches(6.8), Inches(1.65), Inches(5.8), Inches(0.4),
    "正常 vs 缺陷 异常分数分布", size=16, bold=True, color=C_NAVY)
img(s, os.path.join(RESULT_DIR, "perception", "score_distribution.png"),
    Inches(6.9), Inches(2.15), width=Inches(5.6))

# 底部：数据/算力说明
rrect(s, Inches(0.6), Inches(6.3), Inches(12.1), Inches(0.7), C_BLUE_BG)
txt(s, Inches(0.8), Inches(6.38), Inches(11.7), Inches(0.55),
    "数据：MVTec-AD公开数据集（115张测试=93缺陷+22正常）  |  算力：感知层CPU单图2-3秒，认知层CPU异步20-30秒  |  无需企业私有数据",
    size=12, color=C_DARK, anchor=MSO_ANCHOR.MIDDLE)

# ================================================================
# 第11页：先导实验（4张实验图拼版）
# ================================================================
s = blank()
set_bg(s, C_BG)
page_header(s, "五、先导实验 · 实测结果", "核心技术路线可行性已验证，可进入中期开发", 11)

# 2x2 实验图拼版
experiments = [
    ("感知层：高置信快链路", "high_conf_sample.png", "perception", C_GREEN,
     "异常分数0.42 ≥ 阈值0.22，直接输出缺陷定位"),
    ("感知层：低置信触发慢链路", "low_conf_sample.png", "perception", C_ORANGE,
     "异常分数≈阈值，触发VLM+RAG深度分析"),
    ("认知层：RAG标准检索", "retrieval_result.png", "cognition", C_BLUE,
     "TF-IDF检索匹配相关标准条款，带引用ID"),
    ("认知层：完整流程演示", "cognition_demo.png", "cognition", RGBColor(0x6A,0x1B,0x9A),
     "缺陷描述→标准匹配→引用校验→结构化输出"),
]
for i, (title, fname, sub, c, desc) in enumerate(experiments):
    row = i // 2; col = i % 2
    l = Inches(0.6 + col * 6.15)
    t = Inches(1.4 + row * 2.85)
    rrect(s, l, t, Inches(5.95), Inches(2.65), C_WHITE, C_LINE, 1)
    # 标题条
    rect(s, l, t, Inches(5.95), Inches(0.45), c)
    txt(s, l + Inches(0.15), t + Inches(0.05), Inches(4), Inches(0.35),
        title, size=14, bold=True, color=C_WHITE, anchor=MSO_ANCHOR.MIDDLE)
    # 图片
    img(s, os.path.join(RESULT_DIR, sub, fname),
        l + Inches(0.1), t + Inches(0.55), width=Inches(5.75))
    # 说明
    txt(s, l + Inches(0.15), t + Inches(2.25), Inches(5.65), Inches(0.35),
        desc, size=11, color=C_GRAY)

# ================================================================
# 第12页：评估指标与对比实验
# ================================================================
s = blank()
set_bg(s, C_BG)
page_header(s, "六、评估指标与对比实验", "三层量化体系 + 分阶段对比", 12)

# 左：三层指标
txt(s, Inches(0.6), Inches(1.4), Inches(5), Inches(0.4),
    "三层量化评估体系", size=18, bold=True, color=C_NAVY)

metric_layers = [
    ("检测层", "AUROC · 最佳F1 · 定位IoU\n误报率 · 漏报率", C_BLUE),
    ("认知层", "标准匹配Top-1准确率 · 引用准确率\n缺陷描述召回率 · 幻觉率", C_GREEN),
    ("系统层", "单图平均处理耗时 · 报告生成耗时\n人工操作步骤减少率", C_ORANGE),
]
for i, (name, metrics, c) in enumerate(metric_layers):
    t = Inches(1.95 + i * 1.5)
    rrect(s, Inches(0.6), t, Inches(5.8), Inches(1.3), C_WHITE, c, 1.5)
    rect(s, Inches(0.6), t, Inches(1.4), Inches(1.3), c)
    txt(s, Inches(0.6), t + Inches(0.35), Inches(1.4), Inches(0.6),
        name, size=18, bold=True, color=C_WHITE, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    txt(s, Inches(2.2), t + Inches(0.2), Inches(4.0), Inches(1.0),
        metrics, size=13, color=C_DARK, anchor=MSO_ANCHOR.MIDDLE)

# 右：对比实验
txt(s, Inches(6.8), Inches(1.4), Inches(5), Inches(0.4),
    "分阶段对比实验", size=18, bold=True, color=C_NAVY)

comparisons = [
    ("中期", "纯WinCLIP vs 双链路", "验证双链路的分析价值", C_BLUE),
    ("中期", "纯VLM vs 双链路", "验证感知前置的精度提升", C_GREEN),
    ("结题", "无门控串行 vs 双链路", "验证门控的效率价值", C_ORANGE),
]
for i, (stage, compare, purpose, c) in enumerate(comparisons):
    t = Inches(1.95 + i * 1.5)
    rrect(s, Inches(6.8), t, Inches(6.0), Inches(1.3), C_WHITE, c, 1.5)
    # 阶段标签
    rrect(s, Inches(6.95), t + Inches(0.2), Inches(0.9), Inches(0.4), c)
    txt(s, Inches(6.95), t + Inches(0.22), Inches(0.9), Inches(0.36),
        stage, size=13, bold=True, color=C_WHITE, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    txt(s, Inches(8.0), t + Inches(0.18), Inches(4.6), Inches(0.4),
        compare, size=15, bold=True, color=C_DARK)
    txt(s, Inches(8.0), t + Inches(0.65), Inches(4.6), Inches(0.4),
        f"目的：{purpose}", size=12, color=C_GRAY)

# 底部说明
rrect(s, Inches(0.6), Inches(6.5), Inches(12.1), Inches(0.55), C_BLUE_BG)
txt(s, Inches(0.8), Inches(6.55), Inches(11.7), Inches(0.45),
    "对比维度：纯WinCLIP比检测层+系统层 | 纯VLM比认知层+系统层 | 无门控串行比系统层耗时 — 不同方案分维度比各自优势",
    size=12, color=C_DARK, anchor=MSO_ANCHOR.MIDDLE)

# ================================================================
# 第13页：进度计划与团队分工
# ================================================================
s = blank()
set_bg(s, C_BG)
page_header(s, "六、进度计划与团队分工", "12周三阶段交付，严格对应课程节点", 13)

# 上：时间线
phases = [
    ("第4周", "开题答辩", "最小Demo + PPT + 报告", "15%", C_BLUE),
    ("第9-10周", "中期检查", "核心功能Demo + 对比初测", "25%", C_GREEN),
    ("第16周", "结题答辩", "完整系统 + 报告 + 演示视频", "60%", C_ORANGE),
]
# 时间线
rect(s, Inches(1.5), Inches(2.0), Inches(10.3), Pt(3), C_LINE)
for i, (time, name, deliver, ratio, c) in enumerate(phases):
    cx = Inches(2.0 + i * 4.5)
    # 圆点
    dot = s.shapes.add_shape(MSO_SHAPE.OVAL, cx - Inches(0.15), Inches(1.85), Inches(0.3), Inches(0.3))
    dot.fill.solid(); dot.fill.fore_color.rgb = c; dot.line.fill.background()
    # 卡片
    rrect(s, cx - Inches(1.8), Inches(2.4), Inches(3.6), Inches(2.0), C_WHITE, c, 1.5)
    rect(s, cx - Inches(1.8), Inches(2.4), Inches(3.6), Inches(0.5), c)
    txt(s, cx - Inches(1.7), Inches(2.45), Inches(3.4), Inches(0.4),
        f"{time} · {name}", size=15, bold=True, color=C_WHITE, anchor=MSO_ANCHOR.MIDDLE)
    txt(s, cx - Inches(1.6), Inches(3.05), Inches(3.2), Inches(0.6),
        deliver, size=12, color=C_DARK)
    txt(s, cx - Inches(1.6), Inches(3.7), Inches(3.2), Inches(0.5),
        f"考核占比 {ratio}", size=16, bold=True, color=c)

# 下：团队分工
txt(s, Inches(0.6), Inches(4.8), Inches(5), Inches(0.4),
    "团队分工（3人）", size=18, bold=True, color=C_NAVY)

roles = [
    ("项目负责人", "总架构 · 认知层 · 文档总纂 · 答辩主讲", C_BLUE),
    ("感知算法", "WinCLIP部署 · 多尺度优化 · 指标测试", C_GREEN),
    ("系统支持", "界面开发 · 文档排版 · 演示保障", C_ORANGE),
]
for i, (role, duty, c) in enumerate(roles):
    l = Inches(0.6 + i * 4.15)
    rrect(s, l, Inches(5.35), Inches(3.85), Inches(1.5), C_WHITE, c, 1.5)
    rect(s, l, Inches(5.35), Inches(0.12), Inches(1.5), c)
    txt(s, l + Inches(0.3), Inches(5.5), Inches(3.4), Inches(0.4),
        role, size=17, bold=True, color=c)
    txt(s, l + Inches(0.3), Inches(6.0), Inches(3.4), Inches(0.7),
        duty, size=13, color=C_DARK)

# ================================================================
# 第14页：风险预案
# ================================================================
s = blank()
set_bg(s, C_BG)
page_header(s, "七、风险预案", "每条风险都有明确的降级方案", 14)

risks = [
    ("WinCLIP效果不达预期", "替换PatchCore/Anomalib基线，接口兼容", C_RED),
    ("VLM速度过慢", "切换轻量化模型，或高频场景模板化生成", C_ORANGE),
    ("RAG匹配不准/引用错误", "优化分块策略，人工预置核心标准条目", C_BLUE),
    ("门控阈值不合理", "中位数+1.5×MAD辅助标定，支持人工微调", C_GREEN),
    ("JSON Schema不服从", "正则提取中间层；成功率低直接降级人工选", RGBColor(0x6A,0x1B,0x9A)),
    ("真实样本获取困难", "公开真实工业数据集+模拟场景干扰做跨域验证", C_GRAY),
]
for i, (risk, solution, c) in enumerate(risks):
    row = i // 2; col = i % 2
    l = Inches(0.6 + col * 6.15)
    t = Inches(1.5 + row * 1.7)
    rrect(s, l, t, Inches(5.95), Inches(1.5), C_WHITE, c, 1.5)
    # 风险标签
    rrect(s, l + Inches(0.2), t + Inches(0.2), Inches(1.0), Inches(0.35), c)
    txt(s, l + Inches(0.2), t + Inches(0.22), Inches(1.0), Inches(0.31),
        "风险", size=12, bold=True, color=C_WHITE, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    txt(s, l + Inches(1.35), t + Inches(0.2), Inches(4.4), Inches(0.4),
        risk, size=15, bold=True, color=C_DARK)
    # 应对
    txt(s, l + Inches(0.25), t + Inches(0.75), Inches(5.5), Inches(0.6),
        f"→ {solution}", size=13, color=C_GRAY)

# 底部
rrect(s, Inches(0.6), Inches(6.6), Inches(12.1), Inches(0.55), C_NAVY)
txt(s, Inches(0.6), Inches(6.65), Inches(12.1), Inches(0.45),
    "核心原则：系统永远可用 — 任何模块失效都有降级方案，不影响基本质检辅助功能",
    size=14, bold=True, color=C_ORANGE, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)

# ================================================================
# 第15页：致谢
# ================================================================
s = blank()
set_bg(s, C_NAVY)
# 装饰
rect(s, 0, Inches(3.2), SW, Pt(2), C_ORANGE)
rect(s, Inches(5.5), Inches(2.8), Inches(2.3), Pt(3), C_ORANGE)

txt(s, Inches(1), Inches(1.5), Inches(11.3), Inches(1.0),
    "感谢各位老师聆听", size=44, bold=True, color=C_WHITE, align=PP_ALIGN.CENTER)
txt(s, Inches(1), Inches(2.5), Inches(11.3), Inches(0.5),
    "恳请批评指正", size=24, color=C_BLUE_L, align=PP_ALIGN.CENTER)

txt(s, Inches(1), Inches(4.0), Inches(11.3), Inches(0.5),
    "面向金属表面缺陷的双链路门控智能质检辅助系统",
    size=18, color=C_WHITE, align=PP_ALIGN.CENTER)
txt(s, Inches(1), Inches(4.6), Inches(11.3), Inches(0.4),
    "山东大学人工智能学院  ·  2026.09",
    size=14, color=C_GRAY_L, align=PP_ALIGN.CENTER)

# 保存
prs.save(OUT_PATH)
print(f"PPT v2 已生成: {OUT_PATH}")
print(f"总页数: {len(prs.slides)}")
